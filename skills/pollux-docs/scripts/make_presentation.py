#!/usr/bin/env python3
"""Genera la presentación comercial PolluxData en PPTX 16:9.

Uso: make_presentation.py <presentacion.json> [salida.pptx]
"""
import json
import sys
from datetime import date, timedelta
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[3]
ASSETS = ROOT / "skills" / "pollux-brand" / "assets"
PAL = json.loads((ASSETS / "palette.json").read_text())

NAVY = RGBColor(0x06, 0x11, 0x1F)
ORANGE = RGBColor(0xF0, 0x7A, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF6, 0xF7, 0xFB)
BODY = RGBColor(0x5B, 0x65, 0x73)
MUTED = RGBColor(0xA9, 0xB0, 0xBD)
LIGHT1 = RGBColor(0xC8, 0xD4, 0xE6)
LIGHT2 = RGBColor(0x8B, 0x97, 0xA8)

F_HEAVY = "Manrope ExtraBold"
F_HB = "Manrope Bold"
F_BODY = "Inter"
F_SB = "Inter SemiBold"

MASCOTS = {
    "oracle": "pollux-oracle.png",
    "azure": "pollux-azure.png",
    "aws": "pollux-oracle-aws.png",
}

MES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
       "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

SW, SH = 13.333, 7.5


def fmt(n):
    s = f"{n:,.2f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def human_date(iso):
    d = date.fromisoformat(iso)
    return f"{d.day} de {MES[d.month - 1]} de {d.year}"


def tb(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, font=F_BODY, size=14, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         space_after=6, first=False, spc=None):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if spc:
        r._r.get_or_add_rPr().set("spc", str(spc))
    return p


def rule(slide, x, y, w, color=ORANGE, h=0.045):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                  Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def bg(slide, color):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    return rect


def logo(slide, x, y, w=1.0):
    p = ASSETS / "logo" / "pollux-logo.png"
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w))


def kicker(slide, text, x, y, size=13, color=ORANGE):
    t = tb(slide, x, y, SW - 2 * x - 0.4, 0.4)
    para(t, text, font=F_SB, size=size, color=color, first=True, spc=260)


def heading(slide, num, title):
    t = tb(slide, 0.8, 0.5, SW - 1.6, 0.7)
    p = t.paragraphs[0]
    r = p.add_run()
    r.text = f"{num}  —  "
    r.font.name = F_HEAVY
    r.font.size = Pt(24)
    r.font.color.rgb = ORANGE
    r2 = p.add_run()
    r2.text = title
    r2.font.name = F_HEAVY
    r2.font.size = Pt(24)
    r2.font.color.rgb = NAVY
    rule(slide, 0.8, 1.12, 0.75)


def footer(slide, data):
    t = tb(slide, 0.8, 7.02, 6.0, 0.3)
    para(t, f"PolluxData · Ref {data['ref']}", font=F_BODY, size=8.5, color=MUTED,
         first=True)
    t2 = tb(slide, SW - 4.6, 7.02, 3.8, 0.3)
    para(t2, "polluxdata.com", font=F_BODY, size=8.5, color=MUTED,
         align=PP_ALIGN.RIGHT, first=True)


def bullets(slide, items, x, y, w, size=14, color=BODY, gap=10, dash=True):
    t = tb(slide, x, y, w, SH - y - 0.6)
    for i, it in enumerate(items):
        p = t.paragraphs[0] if i == 0 and not t.paragraphs[0].runs else t.add_paragraph()
        p.space_after = Pt(gap)
        if dash:
            r0 = p.add_run()
            r0.text = "—  "
            r0.font.name = F_SB
            r0.font.size = Pt(size)
            r0.font.bold = True
            r0.font.color.rgb = ORANGE
        r = p.add_run()
        r.text = it
        r.font.name = F_BODY
        r.font.size = Pt(size)
        r.font.color.rgb = color


def style_table(tbl):
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def cell_set(cell, text, font=F_BODY, size=12.5, color=BODY, bold=False,
             align=PP_ALIGN.LEFT, fill=None):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def build(data, out):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    bg(s1, NAVY)
    mimg = ASSETS / "mascot" / MASCOTS.get(data.get("partner"), "pollux-oracle.png")
    if mimg.exists():
        s1.shapes.add_picture(str(mimg), Inches(SW - 7.5), 0, height=Inches(SH))
    logo(s1, 0.8, 0.55)
    kicker(s1, "PROPUESTA COMERCIAL", 0.8, 2.15, size=13)
    t = tb(s1, 0.8, 2.6, 4.9, 2.6)
    para(t, data["project_title"], font=F_HEAVY, size=28, color=WHITE, first=True,
         space_after=0)
    t2 = tb(s1, 0.8, 5.35, 4.9, 1.6)
    para(t2, f"Preparada para {data['client']['company']}", font=F_SB, size=13,
         color=LIGHT1, first=True)
    cl = data["client"]
    para(t2, f"{cl['contact']} — {cl['role']} · {human_date(data['date'])}",
         font=F_BODY, size=10.5, color=LIGHT2)
    para(t2, f"Ref {data['ref']}", font=F_BODY, size=10.5, color=LIGHT2)

    s2 = prs.slides.add_slide(blank)
    heading(s2, "01", "RESUMEN EJECUTIVO")
    bullets(s2, data["executive_points"], 0.8, 1.6, 11.0, size=18, gap=22)
    footer(s2, data)

    s3 = prs.slides.add_slide(blank)
    heading(s3, "02", "CONTEXTO Y OBJETIVO")
    t = tb(s3, 0.8, 1.6, 11.2, 1.2)
    para(t, data["context"], font=F_BODY, size=15, color=BODY, first=True)
    bullets(s3, data["goals"], 0.8, 2.95, 11.0, size=15, gap=12)
    footer(s3, data)

    s4 = prs.slides.add_slide(blank)
    heading(s4, "03", "ALCANCE")
    for x0, label, items in ((0.8, "Incluye", data["scope_in"]),
                             (6.93, "No incluye", data["scope_out"])):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0),
                                   Inches(1.6), Inches(5.6), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = SOFT
        card.line.fill.background()
        card.shadow.inherit = False
        t = tb(s4, x0 + 0.35, 1.95, 4.9, 0.4)
        para(t, label, font=F_HB, size=15, color=NAVY, first=True)
        bullets(s4, items, x0 + 0.35, 2.5, 4.9, size=12.5, gap=9)
    footer(s4, data)

    s5 = prs.slides.add_slide(blank)
    heading(s5, "04", "PLAN DE TRABAJO")
    n_ph = len(data["phases"])
    tbl = s5.shapes.add_table(n_ph + 1, 3, Inches(0.8), Inches(1.7),
                              Inches(11.73), Inches(0.6 + 0.8 * n_ph)).table
    style_table(tbl)
    for i, w in enumerate((1.2, 7.53, 3.0)):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.6)
    for r in range(1, n_ph + 1):
        tbl.rows[r].height = Inches(0.8)
    for c, htxt in enumerate(("Fase", "Entregable", "Duración")):
        cell_set(tbl.cell(0, c), htxt, font=F_HB, size=13, color=WHITE, fill=NAVY,
                 align=PP_ALIGN.RIGHT if c == 2 else PP_ALIGN.LEFT)
    for r, ph in enumerate(data["phases"], 1):
        fill = SOFT if r % 2 == 0 else WHITE
        cell_set(tbl.cell(r, 0), f"Fase {r}", font=F_SB, color=NAVY, fill=fill)
        cell_set(tbl.cell(r, 1), ph["deliverable"], fill=fill)
        cell_set(tbl.cell(r, 2), ph["duration"], align=PP_ALIGN.RIGHT, fill=fill)
    note_y = 1.7 + 0.6 + 0.8 * n_ph + 0.2
    t = tb(s5, 0.8, note_y, 11.2, 0.6)
    para(t, "Cada fase se cierra con aceptación escrita antes de continuar.",
         font=F_BODY, size=10.5, color=MUTED, first=True)
    footer(s5, data)

    s6 = prs.slides.add_slide(blank)
    heading(s6, "05", "INVERSIÓN")
    rows = len(data["pricing"]) + 2
    tbl = s6.shapes.add_table(rows, 6, Inches(0.8), Inches(1.7), Inches(11.73),
                              Inches(0.55 + 0.55 * rows)).table
    style_table(tbl)
    for i, w in enumerate((0.55, 4.98, 1.5, 0.9, 1.9, 1.9)):
        tbl.columns[i].width = Inches(w)
    heads = ("#", "Concepto", "Unidad", "Cant.", "P. unit. (USD)", "Total (USD)")
    for c, htxt in enumerate(heads):
        cell_set(tbl.cell(0, c), htxt, font=F_HB, size=12.5, color=WHITE, fill=NAVY,
                 align=PP_ALIGN.RIGHT if c >= 3 else PP_ALIGN.LEFT)
    for r, it in enumerate(data["pricing"], 1):
        fill = SOFT if r % 2 == 0 else WHITE
        cell_set(tbl.cell(r, 0), str(r), align=PP_ALIGN.RIGHT, fill=fill)
        cell_set(tbl.cell(r, 1), it["item"], fill=fill)
        cell_set(tbl.cell(r, 2), it["unit"], fill=fill)
        cell_set(tbl.cell(r, 3), str(it["qty"]), align=PP_ALIGN.RIGHT, fill=fill)
        cell_set(tbl.cell(r, 4), fmt(it["unit_price"]), align=PP_ALIGN.RIGHT, fill=fill)
        cell_set(tbl.cell(r, 5), fmt(it["qty"] * it["unit_price"]), align=PP_ALIGN.RIGHT,
                 fill=fill)
    last = rows - 1
    sub = sum(i["qty"] * i["unit_price"] for i in data["pricing"])
    tax = round(sub * data["tax_rate"] / 100, 2)
    for c in range(6):
        cell_set(tbl.cell(last, c), "", fill=SOFT)
    cell_set(tbl.cell(last, 1), "TOTAL (IVA incluido)", font=F_HEAVY, size=13.5,
             color=ORANGE, fill=SOFT)
    cell_set(tbl.cell(last, 5), fmt(sub + tax), font=F_HEAVY, size=13.5, color=ORANGE,
             align=PP_ALIGN.RIGHT, fill=SOFT)
    t = tb(s6, 0.8, 5.6, 11.2, 0.8)
    para(t, data["payment_terms"], font=F_BODY, size=11.5, color=BODY, first=True)
    footer(s6, data)

    s7 = prs.slides.add_slide(blank)
    heading(s7, "06", "POR QUÉ POLLUXDATA")
    for i, vp in enumerate(data["value_props"]):
        x0 = 0.8 + i * 4.02
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0),
                                   Inches(2.0), Inches(3.7), Inches(3.4))
        card.fill.solid()
        card.fill.fore_color.rgb = SOFT
        card.line.fill.background()
        card.shadow.inherit = False
        t = tb(s7, x0 + 0.3, 2.35, 3.1, 0.5)
        para(t, f"0{i + 1}", font=F_HEAVY, size=20, color=ORANGE, first=True)
        t2 = tb(s7, x0 + 0.3, 3.05, 3.1, 2.2)
        para(t2, vp, font=F_SB, size=13.5, color=NAVY, first=True)
    footer(s7, data)

    s8 = prs.slides.add_slide(blank)
    bg(s8, NAVY)
    logo(s8, 0.8, 0.55)
    kicker(s8, "SIGUIENTES PASOS", 0.8, 2.1, size=13)
    t = tb(s8, 0.8, 2.55, 7.5, 1.1)
    para(t, "¿Arrancamos?", font=F_HEAVY, size=36, color=WHITE, first=True)
    t2 = tb(s8, 0.8, 3.8, 7.5, 2.4)
    for i, step in enumerate(data["cta_steps"]):
        p = t2.paragraphs[0] if i == 0 else t2.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = f"0{i + 1}  "
        r.font.name = F_HEAVY
        r.font.size = Pt(16)
        r.font.color.rgb = ORANGE
        r2 = p.add_run()
        r2.text = step
        r2.font.name = F_BODY
        r2.font.size = Pt(16)
        r2.font.color.rgb = LIGHT1
    t3 = tb(s8, 8.9, 5.0, 3.6, 2.0)
    para(t3, data["sender"]["name"], font=F_SB, size=13, color=WHITE,
         align=PP_ALIGN.RIGHT, first=True)
    para(t3, data["sender"]["email"], font=F_SB, size=13, color=ORANGE,
         align=PP_ALIGN.RIGHT)
    para(t3, "polluxdata.com", font=F_BODY, size=11, color=LIGHT2,
         align=PP_ALIGN.RIGHT)

    if data.get("demo"):
        for s in prs.slides:
            t = s.shapes.add_textbox(Inches(SW - 4.3), Inches(0.12), Inches(3.7),
                                     Inches(0.3))
            tf = t.text_frame
            tf.word_wrap = False
            para(tf, "MUESTRA · DATOS FICTICIOS", font=F_SB, size=10, color=ORANGE,
                 align=PP_ALIGN.RIGHT, first=True)

    prs.save(str(out))


def main():
    src = Path(sys.argv[1])
    data = json.loads(src.read_text())
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".pptx")
    build(data, out)
    print(out)


if __name__ == "__main__":
    main()
